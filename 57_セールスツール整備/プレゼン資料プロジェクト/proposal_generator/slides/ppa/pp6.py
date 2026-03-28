"""
pp6.py - 発電シミュレーション (Power generation simulation)

Layout: A4 landscape
- Header bar
- KPI cards: 年間発電量, 自家消費量, 自家消費率, 年間CO2削減量
- Monthly generation estimate table (12 months)
- Surplus electricity info
"""
from __future__ import annotations
from pathlib import Path
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches
from proposal_generator.utils import (
    CONTENT_TOP, C_DARK, C_LIGHT_ORANGE, C_ORANGE, C_SUB, C_WHITE,
    C_LIGHT_GRAY, FONT_BLACK, FONT_BODY, HEADER_H, MARGIN, SLIDE_H, SLIDE_W,
    add_footer, add_header_bar, add_rect, add_rounded_rect, add_textbox,
    add_section_header, add_table, fmt_num,
)

TITLE = "発電シミュレーション"

# Typical monthly generation distribution (% of annual, approximate for Japan)
MONTHLY_PCT = [6.5, 7.0, 8.5, 9.5, 10.0, 9.5, 10.5, 10.0, 8.5, 8.0, 6.5, 5.5]
MONTH_NAMES = ["1月", "2月", "3月", "4月", "5月", "6月",
               "7月", "8月", "9月", "10月", "11月", "12月"]


def generate(slide, data: dict, logo_path: Path = None) -> None:
    """Render PP6 (generation simulation) onto a blank slide."""
    add_header_bar(slide, TITLE, logo_path)

    y = CONTENT_TOP + Inches(0.1)

    annual_gen = data.get("annual_gen_kwh")
    self_kwh = data.get("self_consumption_kwh")
    self_pct = data.get("self_consumption_pct")
    surplus_kwh = data.get("surplus_kwh")
    _co2_raw = data.get("co2_annual_t")
    # Guard: co2 may be a descriptive string instead of a number
    try:
        co2_t = float(_co2_raw) if _co2_raw is not None else None
    except (ValueError, TypeError):
        co2_t = None

    # ---- KPI cards (4 across) ----
    kpis = [
        (fmt_num(annual_gen, 0) if annual_gen else "—", "kWh/年", "年間発電量"),
        (fmt_num(self_kwh, 0) if self_kwh else "—", "kWh/年", "自家消費量"),
        (_fmt_pct(self_pct), "%", "自家消費率"),
        (fmt_num(co2_t, 1) if co2_t else "—", "t-CO₂/年", "年間CO₂削減量"),
    ]

    card_cols = 4
    gap = Inches(0.15)
    card_w = (SLIDE_W - MARGIN * 2 - gap * (card_cols - 1)) / card_cols
    card_h = Inches(1.15)

    for i, (number, unit, label) in enumerate(kpis):
        cx = MARGIN + i * (card_w + gap)
        add_rounded_rect(slide, cx, y, card_w, card_h, C_LIGHT_ORANGE)
        add_rect(slide, cx, y, card_w, Inches(0.06), C_ORANGE)
        add_textbox(slide, cx, y + Inches(0.12), card_w, Inches(0.45),
                    number,
                    font_name=FONT_BLACK, font_size_pt=26,
                    font_color=C_ORANGE, bold=True,
                    align=PP_ALIGN.CENTER)
        add_textbox(slide, cx, y + Inches(0.55), card_w, Inches(0.2),
                    unit,
                    font_name=FONT_BODY, font_size_pt=9,
                    font_color=C_SUB, align=PP_ALIGN.CENTER)
        add_textbox(slide, cx + Inches(0.06), y + card_h - Inches(0.3),
                    card_w - Inches(0.12), Inches(0.26),
                    label,
                    font_name=FONT_BODY, font_size_pt=10,
                    font_color=C_DARK, bold=True, align=PP_ALIGN.CENTER)

    y += card_h + Inches(0.2)

    # ---- Monthly generation table ----
    add_section_header(slide, MARGIN, y, SLIDE_W - MARGIN * 2,
                       "月別発電量（推定）", font_size_pt=12)
    y += Inches(0.35)

    # Build monthly values
    monthly_kwh = []
    if annual_gen:
        for pct in MONTHLY_PCT:
            monthly_kwh.append(int(float(annual_gen) * pct / 100))
    else:
        monthly_kwh = ["—"] * 12

    # Table: 2 rows of 6 months each for readability
    row1_header = MONTH_NAMES[:6]
    row1_values = [fmt_num(v, 0) if isinstance(v, (int, float)) else v for v in monthly_kwh[:6]]
    row2_header = MONTH_NAMES[6:]
    row2_values = [fmt_num(v, 0) if isinstance(v, (int, float)) else v for v in monthly_kwh[6:]]

    table_w = SLIDE_W - MARGIN * 2
    col_w = table_w / 6
    col_widths = [col_w] * 6

    rows_top = [row1_header, row1_values]
    add_table(slide, MARGIN, y, table_w, rows_top, col_widths, font_size_pt=9)
    y += Inches(0.28) * 2 + Inches(0.08)

    rows_bottom = [row2_header, row2_values]
    add_table(slide, MARGIN, y, table_w, rows_bottom, col_widths, font_size_pt=9)
    y += Inches(0.28) * 2 + Inches(0.2)

    # ---- Surplus electricity ----
    if surplus_kwh:
        add_section_header(slide, MARGIN, y, SLIDE_W - MARGIN * 2,
                           "余剰電力", font_size_pt=12)
        y += Inches(0.35)

        surplus_price = data.get("surplus_price")
        surplus_text = f"年間余剰電力量：{fmt_num(surplus_kwh, 0)} kWh"
        if surplus_price:
            surplus_text += f"　（売電単価：{fmt_num(surplus_price, 1)} 円/kWh）"
        add_textbox(slide, MARGIN + Inches(0.1), y,
                    SLIDE_W - MARGIN * 2 - Inches(0.1), Inches(0.28),
                    surplus_text,
                    font_name=FONT_BODY, font_size_pt=10,
                    font_color=C_DARK)

    add_footer(slide)


def _fmt_pct(val) -> str:
    """Format a percentage value (may be 0-1 float or 0-100 value)."""
    if val is None:
        return "—"
    try:
        v = float(val)
        # If value is <= 1, assume it's a ratio
        if v <= 1.0:
            v *= 100
        return f"{v:.1f}"
    except (TypeError, ValueError):
        return "—"
